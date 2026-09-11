from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from src.exporter import build_main_excel, build_raw_excel
from src.pipeline import run_pipeline


class FakeSerper:
    def __init__(self, _keys):
        self.diagnostics = {
            "requests": 0, "results": 0, "retries": 0, "key_failovers": 0,
            "key_errors": 0, "malformed": 0, "failed_queries": 0,
            "completed_queries": 0,
        }

    def collect(self, _plan, _progress):
        return []


def portal_result(portal_id, name, *_args, **_kwargs):
    status = {"Sumber": name, "Status": "Berhasil", "Ditemukan": 0, "Lolos Periode": 0, "Gagal Parse Tanggal": 0, "Warning/Error": ""}
    if portal_id != "inside_lombok":
        return [], status
    record = {
        "record_id": "one", "source_type": "portal", "source": name,
        "title": "Proyek konstruksi pembangunan hotel baru mulai dibangun di Mandalika",
        "date": date(2026, 8, 18), "date_raw": "18 Agustus 2026",
        "url": "https://example.com/one", "snippet": "Investasi bangunan tumbuh",
        "article_text": "", "source_context_local": True,
    }
    status.update({"Ditemukan": 1, "Lolos Periode": 1})
    return [record], status


def test_representative_pipeline_long_rows_zero_counts_and_exports(monkeypatch, config):
    monkeypatch.setattr("src.pipeline.SerperClient", FakeSerper)
    monkeypatch.setattr("src.pipeline.crawl_portal", portal_result)
    selected = ["LU.F", "LU.I.1", "EXP.4.a", "LU.B.2"]
    result = run_pipeline(2026, 3, selected, config, [], date(2026, 9, 6))
    assert len(result["raw_df"]) == 1
    assert set(result["selected_df"]["taxonomy_code"]) == {"LU.F", "LU.I.1", "EXP.4.a"}
    assert result["summary"]["Classification Rows Terpilih"] == 3
    coal = result["taxonomy_summary"].set_index("Kode").loc["B.2", "Jumlah Berita"]
    assert coal == 0

    main_book = load_workbook(build_main_excel(result), data_only=True)
    lu_book = load_workbook(build_main_excel(result, dimension="LU"), data_only=True)
    exp_book = load_workbook(build_main_excel(result, dimension="EXP"), data_only=True)
    raw_book = load_workbook(build_raw_excel(result), data_only=True)
    assert main_book.sheetnames == ["Ringkasan", "Berita"]
    assert raw_book.sheetnames == ["Raw"]
    assert main_book["Berita"].max_row == 4
    assert lu_book["Berita"].max_row == 3
    assert exp_book["Berita"].max_row == 2
    assert {row[0].value for row in lu_book["Berita"].iter_rows(min_row=2)} == {"Lapangan Usaha"}
    assert {row[0].value for row in exp_book["Berita"].iter_rows(min_row=2)} == {"Pengeluaran"}
    assert raw_book["Raw"].max_row == 2


def test_source_failure_isolated(monkeypatch, config):
    monkeypatch.setattr("src.pipeline.SerperClient", FakeSerper)

    def one_fails(portal_id, name, *_args, **_kwargs):
        if portal_id == "inside_lombok":
            from src.portal_scrapers import PortalScrapingError
            raise PortalScrapingError("fixture failure")
        return portal_result(portal_id, name, *_args, **_kwargs)

    monkeypatch.setattr("src.pipeline.crawl_portal", one_fails)
    result = run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 6))
    assert result["metadata"]["status"] == "Selesai dengan peringatan"
    assert set(result["source_statuses"]["Status"]) == {"Berhasil", "Gagal"}


def test_zero_news_is_a_valid_completed_run(monkeypatch, config):
    monkeypatch.setattr("src.pipeline.SerperClient", FakeSerper)

    def empty_portal(_portal_id, name, *_args, **_kwargs):
        return [], {
            "Sumber": name, "Status": "Berhasil", "Ditemukan": 0, "Lolos Periode": 0,
            "Gagal Parse Tanggal": 0, "Warning/Error": "",
        }

    monkeypatch.setattr("src.pipeline.crawl_portal", empty_portal)
    result = run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 6))
    assert result["metadata"]["status"] == "Selesai"
    assert result["summary"]["Raw Records"] == 0


def test_global_excluded_record_stays_in_raw_without_classification(monkeypatch, config):
    monkeypatch.setattr("src.pipeline.SerperClient", FakeSerper)

    def excluded_portal(portal_id, name, *_args, **_kwargs):
        status = {
            "Sumber": name, "Status": "Berhasil", "Ditemukan": 0, "Lolos Periode": 0,
            "Gagal Parse Tanggal": 0, "Warning/Error": "",
        }
        if portal_id != "inside_lombok":
            return [], status
        status.update({"Ditemukan": 1, "Lolos Periode": 1})
        return [{
            "record_id": "excluded", "source_type": "portal", "source": name,
            "title": "Kecelakaan bus di Lombok Tengah", "date": date(2026, 8, 18),
            "date_raw": "18 Agustus 2026", "url": "https://example.com/excluded",
            "snippet": "Angkutan darat mengalami kecelakaan", "article_text": "",
            "source_context_local": True,
        }], status

    monkeypatch.setattr("src.pipeline.crawl_portal", excluded_portal)
    result = run_pipeline(2026, 3, ["LU.H.2"], config, [], date(2026, 9, 6))
    assert len(result["raw_df"]) == 1
    assert result["raw_df"].iloc[0]["classification_count"] == 0
    assert result["selected_df"].empty


def test_sources_overlap_and_progress_stays_on_caller_thread(monkeypatch, config):
    from threading import Barrier, get_ident
    barrier = Barrier(4)
    caller = get_ident()
    progress_threads = []
    class WaitingSerper(FakeSerper):
        def collect(self, _plan, progress):
            barrier.wait(timeout=5)
            progress("Serper ready")
            return []
    def portal(portal_id, name, *_args, **kwargs):
        if portal_id in {"inside_lombok", "lombok_post", "radar_mandalika"}:
            barrier.wait(timeout=5)
        kwargs["progress"](name)
        return [], {"Sumber": name, "Status": "Berhasil", "Ditemukan": 0,
                    "Lolos Periode": 0, "Gagal Parse Tanggal": 0, "Warning/Error": ""}
    monkeypatch.setattr("src.pipeline.SerperClient", WaitingSerper)
    monkeypatch.setattr("src.pipeline.crawl_portal", portal)
    result = run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 11),
                          progress=lambda *_: progress_threads.append(get_ident()))
    assert result["metadata"]["status"] == "Selesai"
    assert result["source_statuses"]["Sumber"].tolist() == ["Serper", *config.portals["name"]]
    assert set(progress_threads) == {caller}


def test_only_selected_sources_are_called(monkeypatch, config):
    calls = []
    class TrackingSerper(FakeSerper):
        def collect(self, *args):
            calls.append("serper")
            return []
    def portal(portal_id, name, *args, **kwargs):
        calls.append(portal_id)
        return portal_result(portal_id, name)
    monkeypatch.setattr("src.pipeline.SerperClient", TrackingSerper)
    monkeypatch.setattr("src.pipeline.crawl_portal", portal)
    result = run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 11), selected_sources=["radar_mandalika"])
    assert calls == ["radar_mandalika"]
    assert result["query_diagnostics"]["planned_queries"] == 0
    assert result["source_statuses"]["Sumber"].tolist() == ["Radar Mandalika"]
    calls.clear()
    result = run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 11), selected_sources=["serper"])
    assert calls == ["serper"]
    assert result["metadata"]["selected_sources"] == ["serper"]


def test_invalid_or_empty_sources_are_rejected(config):
    import pytest
    for sources in ([], ["unknown"]):
        with pytest.raises(ValueError, match="sumber|Sumber"):
            run_pipeline(2026, 3, ["LU.F"], config, [], date(2026, 9, 11), selected_sources=sources)
