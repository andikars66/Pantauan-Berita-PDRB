from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int) -> None:
    for cell in sheet[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT


def _autosize(sheet, maximum: int = 60) -> None:
    for index, column in enumerate(sheet.columns, start=1):
        length = max((len(str(cell.value)) if cell.value is not None else 0 for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(length + 2, 10), maximum)


def _write_table(sheet, start_row: int, headers: list[str], rows: Iterable[Iterable[Any]]) -> int:
    for column, value in enumerate(headers, start=1):
        sheet.cell(start_row, column, value)
    _style_header(sheet, start_row)
    row_number = start_row + 1
    for values in rows:
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        row_number += 1
    if row_number > start_row + 1:
        sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{row_number - 1}"
    return row_number


def build_main_excel(result: dict[str, Any], dimension: str | None = None) -> BytesIO:
    if dimension not in (None, "LU", "EXP"):
        raise ValueError("dimension must be 'LU', 'EXP', or None")

    dimension_name = {"LU": "Lapangan Usaha", "EXP": "Pengeluaran"}.get(dimension)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Ringkasan"
    metadata = result["metadata"]
    summary_sheet.append(["Ringkasan Pantauan Berita PDRB"])
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append(["Periode", f'{metadata["year"]} T{metadata["quarter"]}'])
    summary_sheet.append(["Rentang", f'{metadata["start_date"]:%d-%m-%Y} s.d. {metadata["end_date"]:%d-%m-%Y}'])
    run_at = metadata["run_at"]
    if getattr(run_at, "tzinfo", None) is not None:
        run_at = run_at.replace(tzinfo=None)
    summary_sheet.append(["Waktu Run", run_at])
    summary_sheet["B4"].number_format = "dd-mm-yyyy hh:mm"
    summary_sheet.append(["Status", metadata["status"]])
    summary_sheet.append([])
    row = _write_table(
        summary_sheet, 7, ["Metric", "Nilai"], result["summary"].items(),
    )
    row += 1
    taxonomy = result["taxonomy_summary"]
    if dimension_name:
        taxonomy = taxonomy[taxonomy["Dimensi"] == dimension_name]
    row = _write_table(
        summary_sheet, row, list(taxonomy.columns), taxonomy.itertuples(index=False, name=None),
    )
    row += 1
    statuses = result["source_statuses"]
    _write_table(summary_sheet, row, list(statuses.columns), statuses.itertuples(index=False, name=None))
    summary_sheet.freeze_panes = "A7"
    _autosize(summary_sheet)

    news_sheet = workbook.create_sheet("Berita")
    headers = ["Dimensi", "Kode", "Sektor/Subsektor", "Judul Berita", "Tanggal", "Pengaruh", "Tautan", "Sumber"]
    rows = []
    selected = result["selected_df"]
    if dimension:
        selected = selected[selected["dimension"] == dimension]
    for item in selected.itertuples():
        rows.append((
            "Lapangan Usaha" if item.dimension == "LU" else "Pengeluaran",
            item.taxonomy_code.split(".", 1)[1], item.label, item.title,
            item.date, item.impact, item.url, item.source,
        ))
    _write_table(news_sheet, 1, headers, rows)
    news_sheet.freeze_panes = "A2"
    for cell in news_sheet["E"][1:]:
        cell.number_format = "dd-mm-yyyy"
    _autosize(news_sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_raw_excel(result: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Raw"
    headers = [
        "Sektor/Subsektor", "Judul Berita", "Tanggal", "Pengaruh", "Tautan", "Sumber",
        "Source Type", "Snippet", "Classification Count",
    ]
    rows = [
        (
            item.classifications, item.title, item.date, item.impacts, item.url, item.source,
            item.source_type, item.snippet, item.classification_count,
        )
        for item in result["raw_df"].itertuples()
    ]
    _write_table(sheet, 1, headers, rows)
    sheet.freeze_panes = "A2"
    for cell in sheet["C"][1:]:
        cell.number_format = "dd-mm-yyyy"
    _autosize(sheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
